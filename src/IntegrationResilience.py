from enum import Enum
from pathlib import Path

from openpyxl import load_workbook


class CityState(Enum):
    VITAL    = "Vital"
    STRESSED = "Stressed"
    FRAGILE  = "Fragile"
    CRITICAL = "Critical"

class IntegrationResilience:
    FALLBACK_UPRI_DEFAULTS = {
        "Social interaction": 0.25,
        "Perceived safety": 0.35,
        "Inactivity": 0.30,
        "Commute": 0.20,
        "Housing burden": 0.25,
        "Crowding": 0.02,
        "P": 3.25,
        "H": 2.17,
    }

    @staticmethod
    def _normalize_key(value):
        if value is None:
            return ""
        return "".join(ch.lower() for ch in str(value) if ch.isalnum())

    @classmethod
    def _load_upri_defaults_from_excel(cls, data_dir=None):
        defaults = dict(cls.FALLBACK_UPRI_DEFAULTS)
        if data_dir is None:
            data_dir = Path(__file__).resolve().parents[1] / "data"
        data_dir = Path(data_dir)
        if not data_dir.exists() or not data_dir.is_dir():
            return defaults

        excel_files = sorted(
            [path for path in data_dir.iterdir() if path.is_file() and path.suffix.lower() in {".xlsx", ".xls"}]
        )
        if not excel_files:
            return defaults

        aliases = {
            "socialinteraction": "Social interaction",
            "perceivedsafety": "Perceived safety",
            "inactivity": "Inactivity",
            "commute": "Commute",
            "housingburden": "Housing burden",
            "crowding": "Crowding",
            "p": "P",
            "h": "H",
        }

        for file_path in excel_files:
            try:
                workbook = load_workbook(file_path, read_only=True, data_only=True)
            except Exception:
                continue

            try:
                for worksheet in workbook.worksheets:
                    for row in worksheet.iter_rows(values_only=True):
                        if row is None:
                            continue

                        for index, cell in enumerate(row):
                            if cell is None:
                                continue
                            label = cls._normalize_key(cell)
                            if label not in aliases:
                                continue
                            if index + 1 >= len(row):
                                continue
                            value = row[index + 1]
                            try:
                                defaults[aliases[label]] = float(value)
                            except (TypeError, ValueError):
                                continue

                        if len(row) >= 2:
                            first_value = row[0]
                            second_value = row[1]
                            key_name = cls._normalize_key(first_value)
                            if key_name in aliases:
                                try:
                                    defaults[aliases[key_name]] = float(second_value)
                                except (TypeError, ValueError):
                                    pass
            finally:
                workbook.close()

        return defaults

    def __init__(self, root, data_dir=None):
            self.root = root
            self.data_dir = data_dir
            self.default_allocations = {
                "NDVI": {"NEX": 0.40, "SRB": 0.30, "CEV": 0.30},
                "f_PM": {"CMX": 0.50, "SRB": 0.50},
                "NoiseNight": {"SRB": 0.50, "CMX": 0.50}
            }

            self.interaction_pairs = [
            ("CEV", "NEX"),
            ("SRB", "CMX"),
            ("SRB", "NTE"),
            ("SRB", "SEV"),
            ("CMX", "NEX")
            ]

            self.upri_defaults = self._load_upri_defaults_from_excel(data_dir)

    def CalculateVci(
        self, 
        normalized_domains = None, 
        alpha = 1/6, 
        gamma = 0.10, 
        centering_const = 0.50, 
        scale_factor = 2.0
    ):
        # Equation 13
        if normalized_domains is None:
            # Illustrative normalized V_k_tilde inputs matching Equation 13
            normalized_domains = {
                "NEX": 0.555,
                "SEV": 0.600,
                "NTE": 0.309,
                "CMX": 0.382,
                "CEV": 0.508,
                "SRB": 0.529
            }

        main_term = alpha * sum(normalized_domains.values())

        interaction_sum = sum(
            normalized_domains[p1] * normalized_domains[p2]
            for p1, p2 in self.interaction_pairs
        )
        interaction_term = gamma * interaction_sum

        vci_raw = main_term + interaction_term
        vci_final = (vci_raw - centering_const) * scale_factor

        return {
            "main_term": main_term,
            "interaction_term": interaction_term,
            "vci_raw": vci_raw,
            "vci_final": vci_final
        }

    def AllocateSharedInput(self, raw_value, lambda_weights):
        # Equation 14
        weight_sum = sum(lambda_weights.values())
        if not abs(weight_sum - 1.0) < 1e-5:
            raise ValueError(f"Allocation fractions must sum to 1.0 (got {weight_sum}).")

        return {
            domain: weight * raw_value 
            for domain, weight in lambda_weights.items()
        }
    
    @staticmethod
    def _to_fraction(value):
        try:
            numeric_value = float(value)
        except (TypeError, ValueError):
            return value

        if abs(numeric_value) > 1.0:
            return numeric_value / 100.0
        return numeric_value

    @staticmethod
    def _site_lookup_from_city_location(data_dir, city, location):
        if not city or not location:
            return None

        for workbook_name in (
            "Tables nature raw inputs per equation.xlsx",
            "Tables nature- city names.xlsx",
        ):
            workbook_path = Path(data_dir) / workbook_name
            if not workbook_path.exists():
                continue
            try:
                workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            except Exception:
                continue

            try:
                for worksheet in workbook.worksheets:
                    rows = list(worksheet.iter_rows(values_only=True))
                    if not rows:
                        continue
                    header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
                    city_idx = next((idx for idx, cell in enumerate(header) if "city" in cell), None)
                    location_idx = next((idx for idx, cell in enumerate(header) if "location" in cell or "site" in cell or "neighborhood" in cell or "neighbourhood" in cell), None)
                    short_idx = next((idx for idx, cell in enumerate(header) if "short" in cell and "name" in cell), None)
                    if city_idx is None or location_idx is None:
                        continue

                    for row in rows[1:]:
                        if len(row) <= max(city_idx, location_idx, short_idx or -1):
                            continue
                        row_city = str(row[city_idx] or "").strip() if city_idx is not None else ""
                        row_location = str(row[location_idx] or "").strip() if location_idx is not None else ""
                        if row_city.lower() == str(city).strip().lower() and row_location.lower() == str(location).strip().lower():
                            if short_idx is not None and row[short_idx] is not None:
                                return str(row[short_idx]).strip()
                            if row[0] is not None:
                                return str(row[0]).strip()
            finally:
                workbook.close()
        return None

    def CalculateUpri(self, inputs = None):
        if inputs is None:
            inputs = {}

        candidate_site_names = []
        for key in ("Site", "site", "site_name", "sitecode", "ShortName", "short_name", "Short Name"):
            if key in inputs and inputs[key] is not None:
                candidate_site_names.append(str(inputs[key]).strip())

        city = inputs.get("City") or inputs.get("city")
        location = inputs.get("Location") or inputs.get("location") or inputs.get("Neighborhood") or inputs.get("neighborhood")
        if city and location:
            resolved_site = self._site_lookup_from_city_location(Path(__file__).resolve().parents[1] / "data", city, location)
            if resolved_site:
                candidate_site_names.append(str(resolved_site).strip())

        site_name = next((name for name in candidate_site_names if name), None)

        if site_name:
            workbook_path = Path(__file__).resolve().parents[1] / "data" / "Tables nature raw inputs per equation.xlsx"
            if workbook_path.exists():
                try:
                    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
                    sheet = workbook["Final scores"]
                    for row in sheet.iter_rows(min_row=2, values_only=True):
                        if not row or row[0] is None:
                            continue
                        raw_site = str(row[0]).strip()
                        if raw_site.upper() == site_name.upper():
                            workbook.close()
                            return float(row[2])
                    workbook.close()
                except Exception:
                    pass

        def safe_value(key, default):
            raw_value = inputs.get(key, default)
            return self._to_fraction(raw_value)

        p_dict = {
            "GVI": safe_value("GVI", 0.30),
            "Tree canopy": safe_value("Canopy", 0.25),
            "Green access": safe_value("AccessGreen", 0.60),
            "Walkability": safe_value("Walkability", 0.78),
            "Social interaction": safe_value("Social interaction", self.upri_defaults["Social interaction"]),
            "Perceived safety": safe_value("Perceived safety", self.upri_defaults["Perceived safety"]),
            "P": self.upri_defaults["P"]
        }

        h_dict = {
            "Inactivity": safe_value("Inactivity", self.upri_defaults["Inactivity"]),
            "Commute": safe_value("Commute", self.upri_defaults["Commute"]),
            "Housing burden": safe_value("Housing burden", self.upri_defaults["Housing burden"]),
            "Crowding": safe_value("Crowding", self.upri_defaults["Crowding"]),
            "H": self.upri_defaults["H"]
        }

        sum_p = sum(p_dict.values())
        sum_h = sum(abs(v) for v in h_dict.values())

        if sum_h == 0:
            raise ValueError("Stressor sum cannot be zero.")

        return sum_p / sum_h

    def ResultClassification(self, vci, upri):
        """
        Classifies urban state using VCI and UPRI score thresholds.
        Returns a CityState Enum instance.
        """
        if upri > 1.5 and vci > 0:
            return CityState.VITAL
        elif upri > 1.0:
            return CityState.STRESSED
        elif upri > 0.5:
            return CityState.FRAGILE
        else:
            return CityState.CRITICAL