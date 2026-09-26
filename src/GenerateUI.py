import re
from pathlib import Path
import tkinter as tk
from tkinter import messagebox, ttk

from openpyxl import load_workbook


class GenerateUI:
    PERCENTAGE_INPUT_KEYS = {
        "GVI",
        "Canopy",
        "LightNight",
        "HumidNight",
        "Walkability",
        "Social interaction",
        "Perceived safety",
        "Inactivity",
        "Commute",
        "Housing burden",
        "Crowding",
    }

    @staticmethod
    def normalize_percentage_input(key, value):
        if key not in GenerateUI.PERCENTAGE_INPUT_KEYS:
            return float(value)

        try:
            numeric_value = float(value)
        except (TypeError, ValueError):
            return value

        if abs(numeric_value) > 1.0:
            return numeric_value / 100.0
        return numeric_value

    DEFAULT_INPUT_CONFIG = {
        "PM2.5": ("PM2.5 (μg m⁻³)", "10.0"),
        "T_day": ("Daytime heat T (°C)", "33.0"),
        "T_night": ("Nighttime heat T (°C)", "26.0"),
        "NDVI": ("NDVI (Landscape vegetation)", "0.45"),
        "GVI": ("GVI (Street-level greenery)", "0.30"),
        "AccessGreen": ("Green-space access", "0.60"),
        "Canopy": ("Tree canopy", "0.25"),
        "Blue": ("Blue-space proximity", "0.20"),
        "Micro": ("Restorative micro-space", "0.40"),
        "Sidewalk": ("Sidewalk contiguity", "0.70"),
        "Walkability": ("Walkability", "78.0"),
        "NoiseNight": ("Night noise dB(A)", "55.0"),
        "Daylight": ("Daytime light Z", "0.30"),
        "LightNight": ("Night light Z", "0.60"),
        "HumidNight": ("Night humidity Z", "0.40"),
    }

    @staticmethod
    def get_data_directory(base_dir=None):
        if base_dir is None:
            base_dir = Path(__file__).resolve().parents[1]
        base_dir = Path(base_dir)

        if base_dir.name.lower() in {"data", "datadir"} and base_dir.exists() and base_dir.is_dir():
            return base_dir

        for candidate in [base_dir / "Data", base_dir / "data"]:
            if candidate.exists() and candidate.is_dir():
                return candidate

        if base_dir.exists() and base_dir.is_dir() and any(path.suffix.lower() in {".xlsx", ".xls"} for path in base_dir.iterdir() if path.is_file()):
            return base_dir

        return base_dir / "Data"

    @staticmethod
    def _normalize_header(value):
        if value is None:
            return ""
        return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())

    @staticmethod
    def _field_alias_map():
        return {
            "pm25": "PM2.5",
            "pm2_5": "PM2.5",
            "pm25ugm3": "PM2.5",
            "pm25gmm3": "PM2.5",
            "tday": "T_day",
            "daytimeheat": "T_day",
            "daytimeheattemp": "T_day",
            "daymaxtempc": "T_day",
            "tnight": "T_night",
            "nighttimeheat": "T_night",
            "nighttimeheattemp": "T_night",
            "nightmintempc": "T_night",
            "ndvi": "NDVI",
            "gvi": "GVI",
            "gvipercent": "GVI",
            "accessgreen": "AccessGreen",
            "greenaccess": "AccessGreen",
            "greenaccess01": "AccessGreen",
            "canopy": "Canopy",
            "treecanopy": "Canopy",
            "treecanopypercent": "Canopy",
            "blue": "Blue",
            "bluespace": "Blue",
            "blueproximity01": "Blue",
            "micro": "Micro",
            "microsite": "Micro",
            "restorativemicrospace01": "Micro",
            "sidewalk": "Sidewalk",
            "sidewalkcontiguity": "Walkability",
            "sidewalkcontiguity01": "Walkability",
            "walkability": "Walkability",
            "safety": "Perceived safety",
            "perceivedsafety": "Perceived safety",
            "socialinteraction": "Social interaction",
            "socialinteraction0100": "Social interaction",
            "inactivity": "Inactivity",
            "inactivitypercent": "Inactivity",
            "commute": "Commute",
            "commutemin": "Commute",
            "housingburden": "Housing burden",
            "housingburdenpercent": "Housing burden",
            "crowding": "Crowding",
            "crowding01": "Crowding",
            "noisenight": "NoiseNight",
            "nightnoiselnightdb": "NoiseNight",
            "nightnoise": "NoiseNight",
            "daylight": "Daylight",
            "daylight01": "Daylight",
            "lightnight": "LightNight",
            "nightlight": "LightNight",
            "nightlightnw": "LightNight",
            "nightlightnW": "LightNight",
            "humidnight": "HumidNight",
            "nighthumidity": "HumidNight",
            "nighthumiditypercent": "HumidNight",
        }

    @staticmethod
    def map_header_to_field(header_value):
        key = GenerateUI._normalize_header(header_value)
        alias_map = GenerateUI._field_alias_map()
        if key in alias_map:
            return alias_map[key]
        key_lower = key.lower()
        if "pm25" in key_lower:
            return "PM2.5"
        if "daymax" in key_lower or "daymaxtemp" in key_lower:
            return "T_day"
        if "nightmin" in key_lower or "nightmintemp" in key_lower:
            return "T_night"
        if "daylight" in key_lower:
            return "Daylight"
        if "nightlight" in key_lower:
            return "LightNight"
        if "nightnoise" in key_lower:
            return "NoiseNight"
        if "nighthumidity" in key_lower:
            return "HumidNight"
        if "greenaccess" in key_lower:
            return "AccessGreen"
        if "treecanopy" in key_lower or "canopy" in key_lower:
            return "Canopy"
        if "blue" in key_lower and "proximity" in key_lower:
            return "Blue"
        if "micro" in key_lower and ("space" in key_lower or "micro" in key_lower):
            return "Micro"
        if "sidewalk" in key_lower and "contiguity" in key_lower:
            return "Sidewalk"
        if "walkability" in key_lower:
            return "Walkability"
        if "ndvi" in key_lower:
            return "NDVI"
        if "gvi" in key_lower:
            return "GVI"
        return None

    @staticmethod
    def _find_workbook_path(data_dir, *keywords):
        data_dir = Path(data_dir)
        if not data_dir.exists():
            return None
        for file_path in sorted(data_dir.iterdir()):
            if not file_path.is_file() or file_path.suffix.lower() not in {".xlsx", ".xls"}:
                continue
            name = file_path.name.lower()
            if all(keyword.lower() in name for keyword in keywords):
                return file_path
        return None

    @staticmethod
    def has_data_files(data_dir=None):
        data_dir = GenerateUI.get_data_directory(data_dir)
        raw_city_path = GenerateUI._find_workbook_path(data_dir, "raw", "inputs", "city")
        city_names_path = GenerateUI._find_workbook_path(data_dir, "city", "name")
        return bool(raw_city_path and city_names_path)

    @staticmethod
    def _iter_rows_from_workbook(file_path):
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows(values_only=True):
                    if any(value is not None and str(value).strip() for value in row):
                        yield row
        finally:
            workbook.close()

    @staticmethod
    def load_city_names(data_dir=None):
        data_dir = GenerateUI.get_data_directory(data_dir)
        if not GenerateUI.has_data_files(data_dir):
            return {}

        city_names_path = GenerateUI._find_workbook_path(data_dir, "city", "name")
        if city_names_path is None:
            return {}

        rows = list(GenerateUI._iter_rows_from_workbook(city_names_path))
        if not rows:
            return {}

        header_row = None
        city_idx = None
        location_idx = None
        for row in rows:
            normalized = [GenerateUI._normalize_header(value) for value in row]
            if any("city" in cell for cell in normalized) and (
                any("location" in cell or "site" in cell for cell in normalized)
                or any("neighborhood" in cell or "neighbourhood" in cell or "district" in cell for cell in normalized)
            ):
                header_row = list(row)
                city_idx = next((index for index, value in enumerate(header_row) if "city" in GenerateUI._normalize_header(value)), None)
                location_idx = next((index for index, value in enumerate(header_row) if "location" in GenerateUI._normalize_header(value) or "site" in GenerateUI._normalize_header(value) or "neighborhood" in GenerateUI._normalize_header(value) or "neighbourhood" in GenerateUI._normalize_header(value)), None)
                break

        if header_row is None or city_idx is None or location_idx is None:
            return {}

        city_map = {}
        current_city = None
        for row in rows:
            if row == header_row:
                continue
            if len(row) <= max(city_idx, location_idx):
                continue

            city_value = row[city_idx]
            location_value = row[location_idx]
            city_text = str(city_value).strip() if city_value is not None else ""
            location_text = str(location_value).strip() if location_value is not None else ""
            if city_text.lower() in {"city", "site", "short name", "location"}:
                continue
            if city_text:
                current_city = city_text
            if not location_text or current_city is None:
                continue
            city_map.setdefault(current_city, [])
            if location_text not in city_map[current_city]:
                city_map[current_city].append(location_text)

        ordered_map = {}
        for city_name in sorted(city_map.keys(), key=lambda name: name.lower()):
            ordered_map[city_name] = sorted(city_map[city_name], key=lambda name: name.lower())
        return ordered_map

    @staticmethod
    def load_values_for_selection(data_dir=None, city=None, neighborhood=None):
        data_dir = GenerateUI.get_data_directory(data_dir)
        if not city or not neighborhood or not GenerateUI.has_data_files(data_dir):
            return {}

        city_names_path = GenerateUI._find_workbook_path(data_dir, "city", "name")
        short_name = None
        if city_names_path is not None:
            rows = list(GenerateUI._iter_rows_from_workbook(city_names_path))
            for row in rows:
                normalized = [GenerateUI._normalize_header(value) for value in row]
                if any("city" in cell for cell in normalized) and any("location" in cell or "site" in cell for cell in normalized):
                    city_idx = next((idx for idx, value in enumerate(row) if "city" in GenerateUI._normalize_header(value)), None)
                    location_idx = next((idx for idx, value in enumerate(row) if "location" in GenerateUI._normalize_header(value) or "site" in GenerateUI._normalize_header(value) or "neighborhood" in GenerateUI._normalize_header(value) or "neighbourhood" in GenerateUI._normalize_header(value)), None)
                    short_idx = next((idx for idx, value in enumerate(row) if "short" in GenerateUI._normalize_header(value) and "name" in GenerateUI._normalize_header(value)), None)
                    if city_idx is None or location_idx is None or short_idx is None:
                        continue

                    current_city = None
                    for candidate_row in rows:
                        if len(candidate_row) <= max(city_idx, location_idx, short_idx):
                            continue
                        candidate_city = str(candidate_row[city_idx] or "").strip() if city_idx is not None and candidate_row[city_idx] is not None else ""
                        candidate_location = str(candidate_row[location_idx] or "").strip()
                        if candidate_city:
                            current_city = candidate_city
                        if candidate_location == str(neighborhood).strip() and (current_city == str(city).strip() or (not candidate_city and current_city == str(city).strip())):
                            short_name = str(candidate_row[short_idx] or "").strip()
                            break
                    if short_name:
                        break

        site_files = [
            GenerateUI._find_workbook_path(data_dir, "raw", "inputs", "city"),
            GenerateUI._find_workbook_path(data_dir, "raw", "inputs", "equation"),
        ]

        alias_map = GenerateUI._field_alias_map()
        values = {}
        for file_path in site_files:
            if file_path is None:
                continue
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            try:
                for worksheet in workbook.worksheets:
                    rows = list(worksheet.iter_rows(values_only=True))
                    if not rows:
                        continue
                    header = list(rows[0])
                    normalized_header = [GenerateUI._normalize_header(value) for value in header]
                    site_idx = next((idx for idx, cell in enumerate(normalized_header) if "site" in cell or ("short" in cell and "name" in cell)), None)
                    city_idx = next((idx for idx, cell in enumerate(normalized_header) if "city" in cell), None)
                    neighborhood_idx = next((idx for idx, cell in enumerate(normalized_header) if "neighborhood" in cell or "neighbourhood" in cell or "location" in cell or "district" in cell), None)

                    for row in rows[1:]:
                        if not row or len(row) <= max([idx for idx in [site_idx, city_idx, neighborhood_idx] if idx is not None], default=-1):
                            continue

                        row_site = str(row[site_idx] or "").strip() if site_idx is not None else ""
                        row_city = str(row[city_idx] or "").strip() if city_idx is not None else ""
                        row_neighborhood = str(row[neighborhood_idx] or "").strip() if neighborhood_idx is not None else ""

                        is_site_match = bool(short_name) and site_idx is not None and row_site == str(short_name).strip()
                        is_city_match = bool(city and neighborhood) and city_idx is not None and neighborhood_idx is not None and row_city == str(city).strip() and row_neighborhood == str(neighborhood).strip()

                        if not (is_site_match or is_city_match):
                            continue

                        if site_idx is not None and row_site:
                            values["Site"] = str(row_site).strip()

                        for index, cell_value in enumerate(row):
                            if index >= len(header):
                                continue
                            target = GenerateUI.map_header_to_field(header[index])
                            if target is None or cell_value is None:
                                continue
                            try:
                                values[target] = float(cell_value)
                            except (TypeError, ValueError):
                                pass
            finally:
                workbook.close()

        return values

    def __init__(self, root, vci_calculator=None):
        self.root = root
        self.vci_calculator = vci_calculator
        self.data_dir = self.get_data_directory()
        self.data_available = self.has_data_files(self.data_dir)
        self.city_map = self.load_city_names(self.data_dir) if self.data_available else {}

        self.yTop = 20
        self.xLeft = 50

        self.XREF = 5.0
        self.ST = 5.0
        self.TTHR = 24.0
        self.TMMT = 26

        self.city_var = tk.StringVar()
        self.neighborhood_var = tk.StringVar()
        self.data_city_label = None
        self.data_neighborhood_label = None
        self.city_dropdown = None
        self.neighborhood_dropdown = None
        self.load_values_btn = None
        self.current_site_code = None

        self.calc_btn = tk.Button(root, text="Calculate VCI", command=lambda: self.OnCalculateClicked(),
                                    bg="#0D15ED", fg="white")
        self.calc_btn.place(x=(self.xLeft + 700), y=self.yTop)

        self.constant_label = tk.Label(root, text="Constant Values:", font=("Arial", 10, "bold", "underline"))
        self.constant_label.place(x=self.xLeft, y=self.yTop + 40)
        self.AddConstantsDisplay()

        self.variable_label = tk.Label(root, text="Variable Input:", font=("Arial", 10, "bold", "underline"))
        self.variable_label.place(x=(self.xLeft + 400), y=(self.yTop + 40))
        self.input_entries = {}

        if self.data_available:
            self.AddDataSelectionControls()
        self.AddInputFields()

        self.exit_btn = tk.Button(
            self.root, text="Exit",
            command=self.ConfirmExit,
            bg="#DC3545", fg="white", padx=10, pady=5)
        self.exit_btn.place(relx=1.0, rely=1.0, x=-10, y=-10, anchor="se")

    def AddDataSelectionControls(self):
        city_row_y = self.yTop + 200
        neighborhood_row_y = city_row_y + 35

        self.data_city_label = tk.Label(self.root, text="City:", font=("Arial", 9, "bold"))
        self.data_city_label.place(x=self.xLeft, y=city_row_y)

        city_values = list(self.city_map.keys())
        self.city_dropdown = ttk.Combobox(
            self.root,
            textvariable=self.city_var,
            state="readonly",
            width=25,
            values=city_values,
        )
        self.city_dropdown.place(x=self.xLeft + 90, y=city_row_y - 2)
        self.city_dropdown.bind("<<ComboboxSelected>>", self.OnCitySelected)

        self.data_neighborhood_label = tk.Label(self.root, text="Neighborhood:", font=("Arial", 9, "bold"))
        self.data_neighborhood_label.place(x=self.xLeft, y=neighborhood_row_y)

        self.neighborhood_dropdown = ttk.Combobox(
            self.root,
            textvariable=self.neighborhood_var,
            state="readonly",
            width=25,
        )
        self.neighborhood_dropdown.place(x=self.xLeft + 90, y=neighborhood_row_y - 2)
        self.neighborhood_dropdown.place_forget()
        self.data_neighborhood_label.place_forget()

        if city_values:
            self.city_dropdown.set(city_values[0])
            self.OnCitySelected()

        self.load_values_btn = tk.Button(
            self.root,
            text="Load Values",
            command=self.OnLoadValuesClicked,
            bg="#1F9D55",
            fg="white",
            padx=10,
            pady=5,
        )
        self.load_values_btn.place(x=self.xLeft + 90, y=neighborhood_row_y + 35)

    def OnCitySelected(self, event=None):
        if not self.city_dropdown or not self.neighborhood_dropdown:
            return
        selection_y = self.yTop + 200
        neighborhood_y = selection_y + 35
        selected_city = self.city_var.get()
        neighborhoods = self.city_map.get(selected_city, [])
        if not neighborhoods:
            if self.neighborhood_dropdown is not None:
                self.neighborhood_dropdown.set("")
                self.neighborhood_dropdown.place_forget()
            if self.data_neighborhood_label is not None:
                self.data_neighborhood_label.place_forget()
            return

        self.neighborhood_dropdown.config(values=neighborhoods)
        self.neighborhood_dropdown.set(neighborhoods[0])
        self.neighborhood_dropdown.place(x=self.xLeft + 90, y=neighborhood_y - 2)
        self.data_neighborhood_label.place(x=self.xLeft, y=neighborhood_y)

    def OnLoadValuesClicked(self):
        if not self.data_available:
            return

        city = self.city_var.get().strip()
        neighborhood = self.neighborhood_var.get().strip()
        if not city or not neighborhood:
            messagebox.showwarning("Selection needed", "Please select a city and a neighborhood before loading values.")
            return

        values = self.load_values_for_selection(self.data_dir, city, neighborhood)
        if not values:
            messagebox.showwarning("No data", f"No values found for {city} / {neighborhood}.")
            return

        self.current_site_code = values.get("Site")
        for key, entry in self.input_entries.items():
            if key in values:
                entry.delete(0, tk.END)
                entry.insert(0, str(self.normalize_percentage_input(key, values[key])))

    def OnCalculateClicked(self):
        if self.vci_calculator:
            self.vci_calculator.CalculateVCI()

    def AddConstantsDisplay(self):
        const_x = self.xLeft
        const_y = self.yTop + 40
        constants = [
            ("X_ref", f"{self.XREF}", "μg m⁻³"),
            ("S_t", f"{self.ST}", "°C"),
            ("T_thr", f"{self.TTHR}", "°C"),
            ("Daytime T_MMT", f"{self.TMMT}", "°C")
        ]

        row_height = 25
        for i, (name, value, unit) in enumerate(constants):
            current_y = const_y + 30 + (i * row_height)
            txt = f"{name} = {value} {unit}"
            lbl = tk.Label(
                self.root,
                text=txt,
                font=("Arial", 9, "italic"),
                fg="#333333",
                anchor="w"
            )
            lbl.place(x=const_x, y=current_y)

    def AddInputFields(self):
        vcmd = (self.root.register(self.ValidateFloat), '%P')
        start_y = self.yTop + 80
        row_height = 30

        for i, (key, (label_text, default_val)) in enumerate(self.DEFAULT_INPUT_CONFIG.items()):
            current_y = start_y + 30 + (i * row_height)

            lbl = tk.Label(self.root, text=label_text, font=("Arial", 9), anchor="w")
            lbl.place(x=self.xLeft + 400, y=current_y, width=180)

            entry = tk.Entry(
                self.root,
                font=("Arial", 9),
                validate="key",
                validatecommand=vcmd
            )

            entry.insert(0, str(self.normalize_percentage_input(key, default_val)))
            entry.place(x=self.xLeft + 590, y=current_y, width=70)

            self.input_entries[key] = entry

    def ValidateFloat(self, P):
        if P == "" or P == "-":
            return True
        try:
            float(P)
            return True
        except ValueError:
            return False

    def ConfirmExit(self):
        if messagebox.askyesno("Quit", "Are you sure you want to exit?"):
            try:
                self.root.destroy()
            except:
                pass

    def AddEquationLabels(self):
        if hasattr(self, "eq_labels") and self.eq_labels:
            return

        self.results_label = tk.Label(
            self.root,
            text="Results:",
            font=("Arial", 10, "bold", "underline")
        )
        self.results_label.place(x=(self.xLeft + 800), y=(self.yTop + 40))

        self.eq_labels = {}
        start_y = self.yTop + 80

        lbl = tk.Label(
            self.root,
            text=f"Equation {1}: --",
            font=("Arial", 10),
            anchor="w"
        )
        lbl.place(x=(self.xLeft + 800), y=start_y)
        start_y += 20

        self.eq_labels[1] = lbl

        for i in range(2, 16):
            lbl = tk.Label(
                self.root,
                text=f"Equation {i}: --",
                font=("Arial", 10),
                anchor="w"
            )
            lbl.place(x=(self.xLeft + 800), y=(start_y + (i - 1) * 40))
            self.eq_labels[i] = lbl

        lbl = tk.Label(
            self.root,
            text="Classification: --",
            font=("Arial", 10, "bold", "underline"),
            fg="blue",
            anchor="w"
        )
        lbl.place(x=(self.xLeft + 800), y=(start_y + len(self.eq_labels) * 40))
        self.eq_labels[16] = lbl
